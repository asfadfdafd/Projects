#!/usr/bin/env python3
"""ArchivAI — презентация для сотрудников архива (не для ИТ и не для руководства).

Правила текста: язык повседневной работы архивиста, без терминов «пайплайн»,
«индексация», «интеграция», «конфигурация», «модель», «сервер». Каждый слайд
отвечает на вопрос «что это меняет в моей работе».

Оформление — фирменный стиль JI Studio (brandBook.pdf):
  Deep Navy 000F85 · Steel Blue 5980A6 · Ink 101F20 · Ground F2F2F3 · Paper FFFFFF
  Barlow Condensed (заголовки) / Barlow (текст)

Сборка:  python3 build_archivist.py [out.pptx]
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches as I, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "img")
OUT = sys.argv[1] if len(sys.argv) > 1 else "ArchivAI_for_archivists.pptx"

NAVY = C(0x00, 0x0F, 0x85)
STEEL = C(0x59, 0x80, 0xA6)
STEEL_D = C(0x3C, 0x5C, 0x7A)
INK = C(0x10, 0x1F, 0x20)
GROUND = C(0xF2, 0xF2, 0xF3)
PAPER = C(0xFF, 0xFF, 0xFF)
RULE = C(0xD8, 0xD8, 0xDC)
MUTE = C(0x6B, 0x70, 0x72)
MUTE_ON_NAVY = C(0xB9, 0xC4, 0xD6)
H = "Barlow Condensed"
B = "Barlow"

prs = Presentation()
prs.slide_width, prs.slide_height = I(13.333), I(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 17
_page = {"n": 0}


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, s, x, y, w, h, size=14, color=INK, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, font=B,
         caps_track=False):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = Pt(spacing)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = font
        if caps_track:
            r.font._rPr.set("spc", "120")
    return tb


def bullets(slide, items, x, y, w, h, size=14, color=INK, spacing=20, gap=9,
            marker_color=STEEL):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(spacing)
        p.space_after = Pt(gap)
        r0 = p.add_run()
        r0.text = "— "
        r0.font.size = Pt(size)
        r0.font.color.rgb = marker_color
        r0.font.name = B
        r = p.add_run()
        r.text = it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = B
    return tb


def card(slide, x, y, w, h, fill=PAPER, line=RULE):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def outline(slide, x, y, w, h, line=STEEL):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    sh.fill.background()
    sh.line.color.rgb = line
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def hairline(slide, x, y, w, color=RULE):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), Emu(9525))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def num_box(slide, x, y, d, label, fill=NAVY, color=PAPER, size=13):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(d), I(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = H
    return sh


def pic(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(os.path.join(IMG, name), I(x), I(y),
                                    I(w) if w else None, I(h) if h else None)


def header(slide, dark=False, fig=None):
    _page["n"] += 1
    mark = "mark_light.png" if dark else "mark_dark.png"
    slide.shapes.add_picture(os.path.join(IMG, mark), I(0.6), I(0.31), height=I(0.17))
    text(slide, "JI STUDIO  —  ARCHIVAI", 0.95, 0.3, 6, 0.24, size=9,
         color=MUTE_ON_NAVY if dark else MUTE, font=H, caps_track=True)
    if fig:
        text(slide, fig, 8.4, 0.3, 3.0, 0.24, size=9, color=MUTE_ON_NAVY if dark else MUTE,
             font=H, align=PP_ALIGN.RIGHT, caps_track=True)
    text(slide, f"{_page['n']:02d} / {TOTAL}", 11.6, 0.3, 1.13, 0.24, size=9,
         color=MUTE_ON_NAVY if dark else MUTE, font=H, align=PP_ALIGN.RIGHT, caps_track=True)
    hairline(slide, 0.6, 0.62, 12.13, color=C(0x2A, 0x36, 0x7A) if dark else RULE)


def head(title, kicker=None, dark=False, fig=None):
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY if dark else GROUND)
    header(s, dark=dark, fig=fig)
    if kicker:
        num_box(s, 0.6, 0.95, 0.3, f"{_page['n'] - 1:02d}",
                fill=STEEL if dark else NAVY, size=11)
        text(s, kicker.upper(), 1.05, 0.98, 8, 0.28, size=13, bold=True,
             color=PAPER if dark else INK, font=H, caps_track=True)
    text(s, title.upper(), 0.6, 1.38, 12.13, 0.72, size=30, bold=True,
         color=PAPER if dark else INK, font=H)
    return s


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


# ═════════════ 1. Титул ═════════════
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
_page["n"] += 1
pic(s, "logo_light.png", 0.6, 0.55, w=1.9)
text(s, "APPLIED AI  ·  AUTONOMOUS SYSTEMS", 0.62, 1.05, 6, 0.24, size=9,
     color=MUTE_ON_NAVY, font=H, caps_track=True)
text(s, "ARCHIVAI", 0.6, 2.3, 9, 1.25, size=76, bold=True, color=PAPER, font=H)
text(s, "Помощник архивиста", 0.6, 3.65, 8.4, 0.6, size=26, color=C(0xD5, 0xDC, 0xE8), font=H)
hairline(s, 0.6, 4.5, 5.2, color=STEEL)
text(s, "Программа читает сканы вместо вас, заполняет описание документа,\n"
        "находит дела по смыслу и собирает именной указатель.\n"
        "Решение всегда остаётся за сотрудником.",
     0.6, 4.75, 8.0, 1.2, size=14, color=MUTE_ON_NAVY, spacing=21)
meta = [("ДЛЯ КОГО", "Сотрудники архива"), ("ЧТО ЭТО", "Рабочая программа"),
        ("ГДЕ РАБОТАЕТ", "На компьютерах архива"), ("ЯЗЫКИ", "Русский · қазақша")]
for i, (k, v) in enumerate(meta):
    y = 2.45 + i * 0.92
    text(s, k, 9.35, y, 3.4, 0.22, size=9, color=STEEL, font=H, caps_track=True)
    text(s, v, 9.35, y + 0.26, 3.4, 0.35, size=15, color=PAPER, font=H)
    if i < 3:
        hairline(s, 9.35, y + 0.72, 3.38, color=C(0x2A, 0x36, 0x7A))
notes(s, "Первая фраза: «это не замена вам, это помощник, который делает рутину».")

# ═════════════ 2. День архивиста ═════════════
s = head("Сколько времени уходит на рутину", "Как это происходит сегодня",
         fig="ЗНАКОМАЯ КАРТИНА")
items = [
    ("Перепечатать текст с листа", "Чтобы дело можно было найти поиском, кто-то должен набрать текст руками."),
    ("Заполнить карточку описания", "Вид документа, дата, кто кому пишет, о чём — всё вносится вручную."),
    ("Вспомнить, где встречалась фамилия", "Приходится держать в голове или искать по описям и своим записям."),
    ("Ответить исследователю", "Запрос «всё о переселении 1930-х» — это дни просмотра дел вручную."),
]
for i, (t, d) in enumerate(items):
    y = 2.35 + i * 1.18
    card(s, 0.6, y, 7.3, 1.0)
    num_box(s, 0.6, y, 0.42, f"{i + 1:02d}", size=12)
    text(s, t, 1.25, y + 0.15, 6.4, 0.3, size=15, bold=True, font=H)
    text(s, d, 1.25, y + 0.5, 6.5, 0.45, size=11.5, color=MUTE, spacing=14)
pic(s, "scan_source.png", 8.4, 2.35, w=2.95)
text(s, "Такой лист сегодня читают и перепечатывают вручную",
     8.4, 6.85, 4.0, 0.5, size=9.5, color=MUTE, spacing=12)
notes(s, "Спросить зал: сколько времени уходит на описание одного дела? Дальше — как это меняется.")

# ═════════════ 3. Что программа делает за вас ═════════════
s = head("Что программа делает за вас", "Простыми словами", fig="ШЕСТЬ ШАГОВ")
steps = [
    ("Прочитает", "разберёт машинопись\nи рукописные пометки"),
    ("Опишет", "вид документа, дата,\nот кого и кому"),
    ("Перескажет", "3–4 предложения\nо чём документ"),
    ("Выпишет имена", "люди, организации,\nместа, даты"),
    ("Свяжет дела", "покажет, где ещё\nвстречается человек"),
    ("Спросит вас", "вы проверяете\nи подтверждаете"),
]
for i, (t, d) in enumerate(steps):
    x = 0.6 + i * 2.06
    card(s, x, 2.5, 1.86, 2.35)
    num_box(s, x, 2.5, 0.42, f"{i + 1:02d}", size=12)
    text(s, t, x + 0.12, 3.15, 1.7, 0.3, size=13.5, bold=True, font=H)
    text(s, d, x + 0.16, 3.5, 1.62, 1.0, size=10.5, color=MUTE, spacing=13.5)
    if i < 5:
        text(s, "→", x + 1.86, 3.45, 0.2, 0.3, size=13, color=STEEL, align=PP_ALIGN.CENTER)
card(s, 0.6, 5.3, 12.13, 1.45)
text(s, "Работа с описанием превращается в проверку", 0.95, 5.52, 11.4, 0.32,
     size=16, bold=True, font=H)
text(s, "Вместо того чтобы набирать текст и заполнять поля с нуля, вы открываете готовое описание, "
        "поправляете неточности и подтверждаете. Всё, что вы исправили, программа запоминает.",
     0.95, 5.9, 11.4, 0.75, size=12, color=MUTE, spacing=16)
notes(s, "Ключ: не «автоматизация вместо человека», а «черновик, который вы проверяете».")

# ═════════════ 4. Как попадают сканы ═════════════
s = head("Сканы попадают в программу сами", "Загрузка", fig="ЭКРАН ЗАГРУЗКИ")
card(s, 0.55, 2.3, 7.3, 4.25)
pic(s, "upload.png", 0.7, 2.45, w=7.0)
bullets(s, [
    "Можно просто перетащить сканы мышкой — по одному или сразу дело целиком",
    "Если сканы складывают в общую папку, программа забирает их оттуда сама, каждые несколько минут",
    "Сотня сканов заливается за один раз — не нужно открывать каждый",
    "Один и тот же скан дважды не попадёт: программа узнаёт уже обработанные",
], 8.35, 2.45, 4.4, 4.0, size=13.5)
notes(s, "Здесь важно: менять привычный порядок сканирования не нужно.")

# ═════════════ 5. Список дел ═════════════
s = head("Всё, что обработано — перед глазами", "Список документов", fig="ЭКРАН СПИСКА")
card(s, 0.55, 2.3, 12.23, 3.75)
pic(s, "documents.png", 0.7, 2.45, w=11.93)
cols = [
    ("Видно, что уже готово", "«обрабатывается» · «ждёт проверки» · «проверено»"),
    ("Видно, что не получилось", "плохой скан не потеряется — он помечен отдельно"),
    ("Видно сам лист", "рядом с названием — маленькое изображение скана"),
]
for i, (t, d) in enumerate(cols):
    x = 0.6 + i * 4.15
    text(s, t, x, 6.3, 3.9, 0.3, size=14, bold=True, font=H)
    text(s, d, x, 6.62, 3.9, 0.5, size=11.5, color=MUTE, spacing=14)
notes(s, "Показать фильтр «ждёт проверки» — это рабочая очередь сотрудника.")

# ═════════════ 6. Карточка документа ═════════════
s = head("Слева лист — справа готовое описание", "Карточка документа", fig="ЭКРАН ДОКУМЕНТА")
card(s, 0.55, 2.3, 5.75, 4.75)
pic(s, "doc_detail.png", 0.7, 2.42, h=4.5)
rows = [
    ("Вид документа и дата", "приказ, протокол, переписка — как принято у вас"),
    ("От кого и кому", "с должностью, если она есть в тексте"),
    ("О чём документ", "короткий пересказ на русском — суть, не читая весь лист"),
    ("Имена и названия", "люди · организации · места · даты · количества"),
    ("Весь текст листа", "можно прочитать, скопировать, вставить в справку"),
]
for i, (t, d) in enumerate(rows):
    y = 2.42 + i * 0.93
    hairline(s, 6.75, y - 0.08, 6.0)
    text(s, t, 6.75, y + 0.05, 5.9, 0.3, size=14.5, bold=True, font=H)
    text(s, d, 6.75, y + 0.38, 5.9, 0.45, size=11.5, color=MUTE, spacing=14)
notes(s, "Обратить внимание: это черновик описания. Правится в один клик — следующий слайд.")

# ═════════════ 7. Проверка ═════════════
s = head("Последнее слово всегда за вами", "Проверка", fig="ЭКРАН ПРОВЕРКИ")
card(s, 6.75, 2.3, 6.03, 4.75)
pic(s, "doc_verify.png", 6.88, 2.42, h=4.5)
bullets(s, [
    "Ошиблась в фамилии — нажмите на неё и исправьте, как в обычном документе",
    "Пропустила организацию — добавьте сами, она встанет в общий указатель",
    "Всё верно — нажмите «Проверено», и дело получает отметку сотрудника",
    "Видно, кто и когда правил: это ваша работа, она зафиксирована",
    "Проверенные дела программа больше не трогает при повторной обработке",
], 0.6, 2.42, 5.75, 3.5, size=13.5)
card(s, 0.6, 5.85, 5.75, 1.2, fill=PAPER, line=STEEL)
text(s, "ПРОГРАММА ПРЕДЛАГАЕТ — СОТРУДНИК РЕШАЕТ", 0.85, 6.05, 5.3, 0.3, size=13.5,
     bold=True, color=NAVY, font=H)
text(s, "Ни одна запись не считается верной, пока её не подтвердил человек.",
     0.85, 6.42, 5.3, 0.45, size=11.5, color=STEEL_D, spacing=14)
notes(s, "Главный слайд для снятия тревоги: без сотрудника система ничего не утверждает.")

# ═════════════ 8. Поиск ═════════════
s = head("Поиск понимает смысл, а не только слова", "Поиск по фонду", fig="ЭКРАН ПОИСКА")
card(s, 0.55, 2.3, 7.35, 2.85)
pic(s, "search_results.png", 0.7, 2.42, w=7.05)
card(s, 0.55, 5.35, 7.35, 1.7)
pic(s, "search_recent.png", 0.7, 5.5, w=7.05)
bullets(s, [
    "Спросите «что-нибудь о кулачестве» — найдутся дела, где этих слов нет, но речь об этом",
    "Не нужно угадывать формулировку из описи или помнить точное название",
    "Ищет по русским и казахским текстам, в том числе в старом написании",
    "Ваши прошлые запросы сохраняются — можно вернуться к ним одним нажатием",
], 8.35, 2.45, 4.4, 4.3, size=13.5)
notes(s, "Сравнить с поиском по описи: там нужно знать слово, здесь — только тему.")

# ═════════════ 9. Вопросы к архиву ═════════════
s = head("Можно просто спросить у архива", "Помощник-собеседник", dark=True,
         fig="ЭКРАН ВОПРОСОВ")
pic(s, "chat.png", 0.6, 2.45, w=6.85)
bullets(s, [
    "Спрашиваете обычными словами: «Кто и кому писал о юбилее Абая?»",
    "Ответ собирается только по вашим делам — ничего не берётся из интернета",
    "Рядом с ответом — ссылки на дела; нажали и открыли сам лист",
    "Если в фонде такого нет, программа так и скажет, а не придумает",
    "Удобно для справок исследователям: черновик ответа со ссылками готов сразу",
], 7.9, 2.5, 4.85, 3.6, size=13.5, color=C(0xE8, 0xEC, 0xF4), marker_color=STEEL)
outline(s, 7.9, 6.15, 4.85, 0.85, line=STEEL)
text(s, "Ответ всегда со ссылкой на дело — можно проверить\nкаждое слово по первоисточнику",
     8.15, 6.3, 4.4, 0.6, size=11.5, bold=True, color=PAPER, spacing=15)
notes(s, "Это про справочную работу: время на подбор дел по запросу резко сокращается.")

# ═════════════ 10. Настройка словами ═════════════
s = head("Настроить можно словами", "Настройка", dark=True,
         fig="ЭКРАН НАСТРОЙКИ")
pic(s, "agent.png", 0.6, 2.45, w=6.95)
text(s, "ПРОСТО НАПИШИТЕ, ЧТО НУЖНО:", 7.95, 2.45, 4.8, 0.3, size=13, bold=True,
     color=STEEL, font=H, caps_track=True)
phrases = [
    "«Добавь вид документа „протокол“\nи перепроверь дела»",
    "«Покажи на главной приказы по годам»",
    "«Ищи ближе по смыслу»",
    "«Объедини одинаковые фамилии»",
]
for i, p_ in enumerate(phrases):
    y = 2.9 + i * 0.78
    outline(s, 7.95, y, 4.8, 0.66, line=C(0x2A, 0x36, 0x7A))
    text(s, p_, 8.15, y + 0.1, 4.5, 0.5, size=11, color=C(0xE8, 0xEC, 0xF4), spacing=13)
outline(s, 7.95, 6.1, 4.8, 0.95, line=STEEL)
text(s, "Программа показывает, что именно изменилось, и любую настройку можно вернуть назад "
        "одной кнопкой. Испортить ничего нельзя.",
     8.15, 6.28, 4.45, 0.7, size=11, color=C(0xD5, 0xDC, 0xE8), spacing=14)
notes(s, "Подчеркнуть: не нужно ждать ИТ-отдел, чтобы добавить вид документа под ваш регламент.")

# ═════════════ 11. Именной указатель ═════════════
s = head("Именной указатель собирается сам", "Люди и организации", fig="ЭКРАН СВЯЗЕЙ")
card(s, 0.55, 2.3, 7.7, 4.6)
pic(s, "entities.png", 0.68, 2.42, w=7.44)
bullets(s, [
    "Одна фамилия — одна карточка, даже если в делах она написана по-разному",
    "Видно все дела, где человек упоминается: не нужно листать описи",
    "Схема показывает, кто с кем встречается в одних и тех же делах",
    "Список можно выгрузить таблицей — для описи, отчёта или справочника",
], 8.6, 2.45, 4.15, 4.3, size=13.5)
notes(s, "Для архивиста это самая понятная выгода: указатель, который иначе делается годами.")

# ═════════════ 12. Сводка и хранилище ═════════════
s = head("Сводка по фонду и условия хранения", "Главная страница", fig="ЭКРАН СВОДКИ")
card(s, 0.55, 2.3, 8.4, 4.6)
pic(s, "dashboard.png", 0.68, 2.42, w=8.14)
bullets(s, [
    "Сколько дел обработано, сколько имён и организаций найдено",
    "По каким годам и языкам распределён фонд",
    "Числа считаются из одного места — они не спорят между собой",
    "Тут же температура и влажность в хранилище, если стоят датчики",
    "Отклонение от нормы видно сразу по цвету",
], 9.25, 2.45, 3.5, 4.4, size=12, spacing=17)
notes(s, "Датчики — по запросу клиента; подключается любой прибор, который умеет отправлять показания.")

# ═════════════ 13. Порядок и учёт ═════════════
s = head("Видно, кто и с чем работал", "Порядок и учёт", fig="ЭКРАН УЧЁТА")
card(s, 4.6, 2.3, 8.18, 4.75)
pic(s, "usage.png", 4.72, 2.42, h=4.5)
bullets(s, [
    "У каждого сотрудника свой вход — работа не смешивается",
    "Записано, кто какие дела открывал и что искал",
    "Видно, какие дела спрашивают чаще всего — это материал для планов оцифровки",
    "Ваши правки видны как ваша работа, а не как «что-то поменялось»",
    "Справочники и массовые операции доступны только ответственному сотруднику",
], 0.6, 2.45, 3.75, 4.6, size=12, spacing=17)
notes(s, "Не «слежка», а прозрачность работы и защита от случайных изменений.")

# ═════════════ 14. Что меняется в работе ═════════════
s = head("Что меняется в вашей работе", "Сравнение", fig="БЫЛО / СТАЛО")
rows = [
    ("Описание одного дела", "набрать текст и заполнить поля вручную",
     "проверить готовое описание"),
    ("Поиск по фонду", "вспомнить формулировку из описи", "спросить своими словами"),
    ("Справка исследователю", "просмотр дел вручную, несколько дней",
     "подборка со ссылками сразу"),
    ("Именной указатель", "собирается годами, вручную", "собирается сам, вы уточняете"),
    ("Разные написания фамилии", "остаются разными записями", "сводятся в одну карточку"),
    ("Знание о фонде", "в голове у опытных сотрудников", "остаётся в системе для всех"),
]
y0, rh = 2.5, 0.68
text(s, "ЗАДАЧА", 0.62, y0, 4.0, 0.3, size=11, bold=True, color=MUTE, font=H, caps_track=True)
text(s, "СЕГОДНЯ", 5.0, y0, 3.9, 0.3, size=11, bold=True, color=MUTE, font=H, caps_track=True)
text(s, "С ПРОГРАММОЙ", 9.3, y0, 3.4, 0.3, size=11, bold=True, color=NAVY, font=H, caps_track=True)
hairline(s, 0.6, y0 + 0.34, 12.13, color=INK)
for i, (a, b_, c_) in enumerate(rows):
    y = y0 + 0.5 + i * rh
    if i % 2 == 0:
        card(s, 0.6, y - 0.06, 12.13, rh - 0.06, fill=PAPER, line=None)
    text(s, a, 0.75, y + 0.06, 4.1, 0.45, size=12.5, bold=True, font=H)
    text(s, b_, 5.0, y + 0.06, 4.1, 0.45, size=12, color=MUTE)
    text(s, c_, 9.3, y + 0.06, 3.4, 0.45, size=12, bold=True, color=NAVY)
notes(s, "Самый важный слайд для сотрудников: показывает выгоду в их же задачах.")

# ═════════════ 15. Про ошибки ═════════════
s = head("Программа ошибается — и мы это считаем", "Честно об ошибках",
         fig="КАЧЕСТВО ЧТЕНИЯ")
cd = CategoryChartData()
cd.categories = ["Как было\nсначала", "Настройка\nчтения", "Убрали лишнюю\nобработку",
                 "Лучшее\nсочетание", "Отсев чужих\nбукв"]
cd.add_series("Ошибок при чтении листа", (0.309, 0.296, 0.219, 0.202, 0.198))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, I(0.6), I(2.4), I(7.5), I(3.9), cd)
ch = gf.chart
ch.has_title = False
ch.has_legend = False
plot = ch.plots[0]
plot.gap_width = 70
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = '0%'
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(11)
dl.font.bold = True
dl.font.color.rgb = INK
dl.font.name = H
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = NAVY
ca = ch.category_axis
ca.has_major_gridlines = False
ca.format.line.color.rgb = RULE
ca.tick_labels.font.size = Pt(10)
ca.tick_labels.font.color.rgb = MUTE
ca.tick_labels.font.name = B
va = ch.value_axis
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = RULE
va.minimum_scale = 0.0
va.maximum_scale = 0.36
va.major_unit = 0.1
va.tick_labels.number_format = '0%'
va.tick_labels.number_format_is_linked = False
va.tick_labels.font.size = Pt(9)
va.tick_labels.font.color.rgb = MUTE
va.tick_labels.font.name = B
text(s, "Доля неверно прочитанных букв на вашем листе 1931 года. Ниже — лучше.",
     0.6, 6.4, 7.5, 0.45, size=11, color=MUTE, italic=True, spacing=13)
bullets(s, [
    "Мы взяли ваш лист, набрали текст вручную как эталон и сверяем с ним каждую настройку",
    "Изменение оставляем, только если ошибок стало меньше — не «на глаз»",
    "Ошибки чаще всего в цифрах и в плохо пропечатанных местах — их и проверяет сотрудник",
    "Чем лучше скан, тем меньше ошибок: с хорошим сканером результат заметно выше",
], 8.4, 2.5, 4.35, 3.4, size=12, spacing=16)
card(s, 8.4, 5.75, 4.35, 1.3, fill=PAPER, line=STEEL)
text(s, "ПОЭТОМУ ПРОВЕРКА ОБЯЗАТЕЛЬНА", 8.65, 5.92, 3.9, 0.28, size=12.5, bold=True,
     color=NAVY, font=H)
text(s, "Программа делает черновик, точность растёт, но подпись под описанием — ваша.",
     8.65, 6.26, 3.9, 0.65, size=11, color=STEEL_D, spacing=14)
notes(s, "Честность здесь работает лучше обещаний: архивисты сразу проверяют такие вещи.")

# ═════════════ 16. Частые вопросы ═════════════
s = head("Вопросы, которые задают первыми", "Частые вопросы", fig="ОТВЕТЫ")
faq = [
    ("Меня заменят?", "Нет. Программа не подписывает описание и не отвечает исследователю — "
                      "она готовит черновик, чтобы вы не набирали текст вручную."),
    ("Куда попадут документы?", "Никуда. Всё остаётся на компьютерах архива, ничего не уходит "
                                "в интернет и посторонним."),
    ("Нужно ли учиться?", "Несколько часов. Экраны как в обычной программе: список дел, карточка, "
                          "поиск. Отдельно покажем на ваших делах."),
    ("А если ошибётся?", "Вы исправляете прямо в карточке — это одно нажатие. Исправления "
                         "сохраняются, и дальше программа читает точнее."),
]
for i, (q, a) in enumerate(faq):
    x = 0.6 + (i % 2) * 6.28
    y = 2.45 + (i // 2) * 2.2
    card(s, x, y, 5.85, 1.9)
    num_box(s, x, y, 0.42, "?", size=15)
    text(s, q, x + 0.62, y + 0.12, 5.0, 0.32, size=16, bold=True, font=H)
    text(s, a, x + 0.62, y + 0.55, 5.0, 1.2, size=11.5, color=MUTE, spacing=15)
hairline(s, 0.6, 6.95, 12.13, color=NAVY)
text(s, "Любой вопрос можно задать прямо сейчас — и мы проверим ответ на ваших делах.",
     0.6, 7.05, 12.13, 0.3, size=12.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER, font=H)
notes(s, "Не пропускать этот слайд: именно эти четыре вопроса решают отношение сотрудников.")

# ═════════════ 17. С чего начнём ═════════════
s = head("С чего начнём", "Дальнейшие шаги", dark=True, fig="ПЛАН")
steps = [
    ("ШАГ 1", "Попробуем на ваших делах",
     "Возьмём 200–500 сканов вашего фонда, настроим виды документов под ваши правила описания "
     "и покажем результат на знакомых вам делах."),
    ("ШАГ 2", "Поставим в архиве и научим",
     "Установим на компьютеры архива, заведём сотрудников, проведём короткое обучение "
     "на рабочих местах."),
    ("ШАГ 3", "Будем улучшать вместе",
     "Ваши правки делают чтение точнее. Добавим то, что нужно именно вам: формы справок, "
     "вырезку подписей и печатей, датчики хранилища."),
]
for i, (badge, title, desc) in enumerate(steps):
    y = 2.5 + i * 1.35
    hairline(s, 0.6, y - 0.12, 12.13, color=C(0x2A, 0x36, 0x7A))
    card(s, 0.6, y, 1.5, 0.42, fill=STEEL, line=None)
    text(s, badge, 0.6, y + 0.08, 1.5, 0.28, size=11, bold=True, color=PAPER,
         align=PP_ALIGN.CENTER, font=H, caps_track=True)
    text(s, title, 2.4, y - 0.02, 10.2, 0.35, size=18, bold=True, color=PAPER, font=H)
    text(s, desc, 2.4, y + 0.4, 10.2, 0.75, size=12, color=MUTE_ON_NAVY, spacing=16)
hairline(s, 0.6, 6.55, 12.13, color=STEEL)
text(s, "От вас нужно только одно: показать, как вы описываете дела сегодня.",
     0.6, 6.7, 8.0, 0.35, size=13.5, bold=True, color=PAPER, font=H)
pic(s, "logo_light.png", 11.05, 6.62, w=1.68)
notes(s, "Закрытие: предложить пилот на их делах и попросить показать текущий регламент описания.")

prs.save(OUT)
print("saved:", OUT)
