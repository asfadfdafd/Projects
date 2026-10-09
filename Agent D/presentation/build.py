#!/usr/bin/env python3
"""ArchivAI — продающая презентация для клиента-архива (RU, 16:9).
Собирается python-pptx: python3 build.py [out.pptx]
Скриншоты берутся из ./img (см. README рядом)."""
import os
import sys

from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "img")
OUT = sys.argv[1] if len(sys.argv) > 1 else "ArchivAI_presentation.pptx"

# Палитра — из интерфейса самого продукта
NAVY, NAVY2 = C(0x0F, 0x1B, 0x2D), C(0x1B, 0x2B, 0x45)
ORANGE = C(0xF9, 0x73, 0x16)
INK, SLATE = C(0x0F, 0x17, 0x2A), C(0x64, 0x74, 0x8B)
WHITE = C(0xFF, 0xFF, 0xFF)
GREEN, GREEN_D = C(0x05, 0x96, 0x69), C(0x04, 0x78, 0x57)
LINE = C(0xE2, 0xE8, 0xF0)
BG_SOFT, BG_CARD = C(0xF8, 0xFA, 0xFC), C(0xF1, 0xF5, 0xF9)
MINT, AMBER_BG, AMBER_TX = C(0xEC, 0xFD, 0xF5), C(0xFF, 0xF7, 0xED), C(0x9A, 0x34, 0x12)
MUTED = C(0x94, 0xA3, 0xB8)
ICE = C(0xCB, 0xD5, 0xE1)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = I(13.333), I(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, s, x, y, w, h, size=14, color=INK, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, space_after=0):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = s.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = Pt(spacing)
        if space_after:
            p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = FONT
    return tb


def bullets(slide, items, x, y, w, h, size=15, color=INK, spacing=21, gap=9):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(spacing)
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = "•  " + it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = FONT
    return tb


def card(slide, x, y, w, h, fill=BG_CARD, line=None, radius=0.06):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    sh.text_frame.text = ""
    return sh


def circle(slide, x, y, d, fill, label=None, size=15, color=WHITE):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(d), I(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    if label is not None:
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = color
        r.font.name = FONT
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return sh


def pic(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(os.path.join(IMG, name), I(x), I(y),
                                    I(w) if w else None, I(h) if h else None)


def head(title, kicker=None, dark=False):
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY if dark else WHITE)
    if kicker:
        text(s, kicker.upper(), 0.6, 0.38, 8, 0.3, size=12, bold=True, color=ORANGE)
    text(s, title, 0.6, 0.68 if kicker else 0.5, 12.1, 0.8, size=31, bold=True,
         color=WHITE if dark else INK)
    return s


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


# ─────────────────────────── 1. Титул ───────────────────────────
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
tb = s.shapes.add_textbox(I(0.9), I(2.0), I(8), I(1.2))
tf = tb.text_frame
tf.margin_left = tf.margin_top = 0
p = tf.paragraphs[0]
for t, col in (("Archiv", WHITE), ("AI", ORANGE)):
    r = p.add_run()
    r.text = t
    r.font.size = Pt(58)
    r.font.bold = True
    r.font.color.rgb = col
    r.font.name = FONT
text(s, "Искусственный интеллект для исторического архива", 0.9, 3.25, 10, 0.6,
     size=23, color=ICE)
text(s, "Сканы → распознанный текст, структурированные поля, связи между людьми и организациями,\n"
        "семантический поиск и проверка человеком.\nПолностью на открытых моделях, на вашем оборудовании.",
     0.9, 4.1, 10.4, 1.3, size=15, color=MUTED, spacing=21)
for i, t in enumerate(["Открытый код", "On-premise", "Русский + қазақша", "Без облаков и подписок"]):
    card(s, 0.9 + i * 2.75, 5.75, 2.55, 0.5, NAVY2, line=C(0x33, 0x41, 0x55))
    text(s, t, 0.9 + i * 2.75, 5.87, 2.55, 0.35, size=12, color=ICE, align=PP_ALIGN.CENTER)
notes(s, "Титул. Тезис: это работающая система, а не концепт — дальше живые экраны.")

# ─────────────────────────── 2. Задача ───────────────────────────
s = head("Что мешает работать с фондом сегодня", "Задача")
items = [
    ("Документы не ищутся", "Скан — это картинка. Найти дело можно только по описи и памяти сотрудника."),
    ("Связи не видны", "Кто с кем переписывался, где упоминается человек — собирается вручную, месяцами."),
    ("Ручная опись", "Тип, дата, отправитель, получатель, персоналии — всё вносится руками."),
    ("Данные должны остаться у вас", "Облачные сервисы и передача фонда наружу — неприемлемы."),
]
for i, (t, d) in enumerate(items):
    y = 1.8 + i * 1.24
    card(s, 0.6, y, 7.3, 1.06)
    circle(s, 0.85, y + 0.28, 0.5, ORANGE, str(i + 1))
    text(s, t, 1.55, y + 0.16, 6.1, 0.32, size=15, bold=True)
    text(s, d, 1.55, y + 0.5, 6.2, 0.5, size=12, color=SLATE, spacing=15)
pic(s, "scan_source.png", 8.35, 1.8, w=3.05)
text(s, "Типовой лист фонда: машинопись, правки от руки, печать, пожелтевшая бумага",
     8.35, 6.5, 4.0, 0.6, size=10, color=SLATE, italic=True, spacing=13)
notes(s, "Боли клиента. Последний пункт подводит к on-prem и открытым моделям.")

# ─────────────────────────── 3. Как это работает ───────────────────────────
s = head("От скана до знания — за один проход", "Как это работает")
steps = [
    ("Приём", "загрузка или забор\nиз вашего хранилища"),
    ("Распознавание", "печатный и рукописный\nтекст, RU + KZ"),
    ("Извлечение", "тип, дата, отправитель,\nполучатель, резюме"),
    ("Сущности", "люди, организации,\nместа, даты, количества"),
    ("Индексация", "поиск по смыслу,\nа не по буквам"),
    ("Проверка", "правка человеком\nи статус «проверено»"),
]
for i, (t, d) in enumerate(steps):
    x = 0.6 + i * 2.06
    card(s, x, 2.05, 1.86, 2.45, BG_SOFT, line=LINE)
    circle(s, x + 0.63, 2.25, 0.6, NAVY, str(i + 1), size=16)
    text(s, t, x + 0.08, 3.0, 1.7, 0.32, size=13.5, bold=True, align=PP_ALIGN.CENTER)
    text(s, d, x + 0.08, 3.36, 1.7, 1.0, size=10.5, color=SLATE, align=PP_ALIGN.CENTER, spacing=14)
    if i < 5:
        text(s, "→", x + 1.8, 3.02, 0.32, 0.4, size=17, color=ORANGE, align=PP_ALIGN.CENTER)
card(s, 0.6, 5.0, 12.1, 1.5, BG_CARD)
text(s, "Каждый этап — отдельный видимый процесс", 0.95, 5.2, 11.4, 0.35, size=15, bold=True)
text(s, "Модель распознавания, правила извлечения, набор типов документов можно заменить или дополнить "
        "без переписывания системы. Ничего не «зашито» в чёрный ящик: любой шаг открывается, повторяется "
        "и настраивается под ваш фонд.",
     0.95, 5.56, 11.4, 0.8, size=12.5, color=SLATE, spacing=17)
notes(s, "Ключевая мысль: конвейер прозрачен и расширяем, модели меняются как детали.")

# ─────────────────────────── 4. Приём документов ───────────────────────────
s = head("Загрузка вручную — и забор пачками из хранилища", "Приём документов")
pic(s, "upload.png", 0.6, 2.0, w=7.2)
bullets(s, [
    "Перетащите сканы в окно — или соберите многостраничное дело из нескольких файлов",
    "Кнопка «Забрать из хранилища»: система сама берёт новые файлы — 50–100+ документов за раз",
    "Автоматический забор каждые 5 минут: положили сканы в папку — они появились в архиве обработанными",
    "Повторная заливка не создаёт дублей: уже обработанные файлы пропускаются",
], 8.2, 2.05, 4.5, 4.0)
notes(s, "Связка с их файловым хранилищем: системе нужны только хранилище и база.")

# ─────────────────────────── 5. Архив ───────────────────────────
s = head("Весь фонд — в одной таблице", "Архив документов")
pic(s, "documents.png", 0.6, 1.85, w=12.1)
cols = [
    ("Статусы обработки", "обрабатывается · на проверке · проверено · ошибка"),
    ("Фильтры", "по типу, языку, статусу, дате документа"),
    ("Миниатюры", "виден сам скан, а не только имя файла"),
]
for i, (t, d) in enumerate(cols):
    x = 0.6 + i * 4.15
    text(s, t, x, 5.95, 3.9, 0.32, size=14, bold=True)
    text(s, d, x, 6.3, 3.9, 0.6, size=12, color=SLATE, spacing=16)
notes(s, "Статус обработки виден сразу, ошибки не теряются.")

# ─────────────────────────── 6. Карточка документа ───────────────────────────
s = head("Скан слева — извлечённые данные справа", "Карточка документа")
pic(s, "doc_detail.png", 0.6, 1.85, h=5.1)
rows = [
    ("Тип, дата, язык", "классификация по вашему справочнику типов"),
    ("Отправитель и получатель", "с должностью и ролью, если они есть в тексте"),
    ("Резюме документа", "2–4 предложения на русском: суть дела без чтения листа"),
    ("Сущности по 5 категориям", "люди · организации · места · даты · количества"),
    ("Полный текст", "постранично и целиком, с источником и уверенностью"),
]
for i, (t, d) in enumerate(rows):
    y = 2.0 + i * 1.0
    circle(s, 6.7, y + 0.07, 0.26, ORANGE)
    text(s, t, 7.15, y, 5.55, 0.34, size=15, bold=True)
    text(s, d, 7.15, y + 0.35, 5.55, 0.5, size=12, color=SLATE, spacing=15)
notes(s, "Резюме и поля извлекаются автоматически — экономия часов описания.")

# ─────────────────────────── 7. Человек в контуре ───────────────────────────
s = head("Последнее слово — за архивистом", "Проверка и доверие")
pic(s, "doc_verify.png", 6.8, 1.85, h=5.05)
bullets(s, [
    "Любую сущность можно исправить прямо в карточке: опечатка в фамилии, пропущенная организация",
    "Кнопка «Проверено» переводит документ в доверенный статус",
    "Каждая правка записывается в журнал: кто, что и когда изменил",
    "Проверенные документы защищены: массовая переобработка их не трогает",
    "Исправления накапливаются как эталон — на них система дообучается под ваш фонд",
], 0.6, 2.0, 5.9, 4.3, size=14.5)
card(s, 0.6, 5.75, 5.9, 1.15, MINT)
text(s, "ИИ предлагает — человек утверждает", 0.9, 5.95, 5.3, 0.35, size=15, bold=True, color=GREEN)
text(s, "Ни одна запись не попадает в фонд как «истина» без подтверждения сотрудника.",
     0.9, 6.32, 5.3, 0.5, size=12, color=GREEN_D, spacing=16)
notes(s, "Ответ на главный страх архива: ИИ ничего не портит и не решает за людей.")

# ─────────────────────────── 8. Поиск ───────────────────────────
s = head("Поиск по смыслу, а не по буквам", "Семантический поиск")
pic(s, "search_results.png", 0.6, 1.85, w=7.3)
pic(s, "search_recent.png", 0.6, 4.85, w=7.3)
bullets(s, [
    "Запрос обычными словами: «что-нибудь о кулачестве» находит приказ о ликвидации кулачества",
    "Работает на русском и казахском, в том числе на старых написаниях",
    "У каждого результата — оценка релевантности; слабые совпадения отсекаются",
    "История запросов у каждого сотрудника своя: можно вернуться к прошлому поиску",
], 8.2, 2.05, 4.5, 4.3)
notes(s, "Разница с поиском по подстроке — это клиент просил явно.")

# ─────────────────────────── 9. Ask ArchivAI ───────────────────────────
s = head("Вопрос архиву — ответ со ссылками на дела", "Ассистент-исследователь", dark=True)
pic(s, "chat.png", 0.6, 2.0, w=6.9)
bullets(s, [
    "Обычный вопрос вместо запроса: «Кто и кому писал о юбилейных мероприятиях?»",
    "Ответ составляется только по вашим документам — каждая мысль подкреплена ссылкой",
    "Клик по ссылке открывает документ-первоисточник",
    "Если ответа в фонде нет, ассистент прямо об этом говорит и не выдумывает",
    "Диалог помнит контекст: вопрос можно уточнять дальше",
], 8.0, 2.05, 4.7, 4.2, color=C(0xE2, 0xE8, 0xF0), size=14.5)
card(s, 8.0, 6.15, 4.7, 0.8, NAVY2)
text(s, "У показанной ранее системы этого нет:\nеё поиск возвращает только список документов",
     8.25, 6.3, 4.3, 0.55, size=11.5, bold=True, color=ORANGE, spacing=15)
notes(s, "Дифференциатор №1. У конкурента это только пожелание клиента, у нас работает.")

# ─────────────────────────── 10. Агент ───────────────────────────
s = head("Настройка системы обычными словами", "Агент · флагман", dark=True)
pic(s, "agent.png", 0.6, 2.0, w=7.0)
text(s, "Скажите — и система перестроится:", 8.1, 2.0, 4.6, 0.35, size=15, bold=True, color=WHITE)
phrases = [
    "«Добавь тип документа „протокол“\nи переклассифицируй архив»",
    "«Добавь виджет: динамика приказов по годам»",
    "«Сделай поиск более смысловым»",
    "«Слей дубликаты персоналий»",
]
for i, p_ in enumerate(phrases):
    y = 2.45 + i * 0.75
    card(s, 8.1, y, 4.6, 0.62, NAVY2, line=C(0x33, 0x41, 0x55))
    text(s, p_, 8.28, y + 0.09, 4.3, 0.5, size=11, color=ICE, spacing=13)
card(s, 8.1, 5.5, 4.6, 1.4, NAVY2)
text(s, "Безопасно по построению", 8.35, 5.62, 4.1, 0.3, size=13, bold=True, color=ORANGE)
text(s, "Модель не имеет доступа к базе: она лишь предлагает изменение. База проверяет его по правилам, "
        "применяет и записывает в журнал. Любую настройку можно откатить одной кнопкой.",
     8.35, 5.95, 4.1, 0.9, size=11, color=ICE, spacing=14)
notes(s, "Дифференциатор №2. Конкурент это только обещал — «поэтому продукт и назван Agent D».")

# ─────────────────────────── 11. Сущности и граф ───────────────────────────
s = head("Люди, организации и места — связаны между делами", "Сущности и связи")
pic(s, "entities.png", 0.6, 1.85, w=7.6)
bullets(s, [
    "Одна персона — одна карточка: разные написания фамилии сводятся автоматически, "
    "старые сохраняются как псевдонимы",
    "Граф показывает, кто встречается в одних и тех же делах — связи, которые вручную не увидеть",
    "От сущности — переход ко всем её документам",
    "Выгрузка списка сущностей в CSV для отчётов и описей",
], 8.5, 2.05, 4.2, 4.4)
notes(s, "Наше преимущество: у конкурента дубликаты «Абай» / «Абай Кунанбаев» живут отдельно.")

# ─────────────────────────── 12. Дашборд + IoT ───────────────────────────
s = head("Состояние фонда — и состояние хранилища", "Аналитика")
pic(s, "dashboard.png", 0.6, 1.85, w=8.3)
bullets(s, [
    "Сколько документов обработано, сколько персоналий и организаций найдено",
    "Распределение по типам, языкам, хронологии фонда",
    "Все цифры считаются из одного источника — они не расходятся между экранами",
    "Датчики хранилища: температура и влажность в реальном времени, с нормами и историей за сутки",
    "Любой ваш датчик шлёт данные на один адрес — виджет появляется сам",
], 9.15, 2.05, 3.65, 4.6, size=12.5, spacing=18)
notes(s, "IoT-виджет — прямой запрос клиента на демо конкурента.")

# ─────────────────────────── 13. Комплаенс ───────────────────────────
s = head("Кто, что и когда смотрел", "Учёт и контроль")
pic(s, "usage.png", 4.7, 1.85, h=5.05)
bullets(s, [
    "Каждый вход — с адресом и временем, включая неудачные попытки",
    "Каждый поиск и каждое открытие документа фиксируются",
    "Видно, какие дела востребованы и кто из сотрудников сейчас работает",
    "Права: администратор ведёт справочники, исследователь работает с фондом",
    "Готово к подключению вашей единой авторизации (SSO)",
], 0.6, 2.05, 3.85, 4.8, size=12.5, spacing=18)
notes(s, "Ответ на «compliance»: журнал, адреса, востребованность документов.")

# ─────────────────────────── 14. Интеграция ───────────────────────────
s = head("Системе нужны только ваше хранилище и ваша база", "Интеграция")
boxes = [
    (0.6, "Ваше файловое\nхранилище", "сканы как есть", BG_CARD, INK, SLATE),
    (4.85, "ArchivAI", "распознавание, извлечение,\nпоиск, проверка", NAVY, WHITE, ICE),
    (9.1, "Ваша база\nи учётная система", "структурированные данные", BG_CARD, INK, SLATE),
]
for x, t, d, fill, col, sub in boxes:
    card(s, x, 2.35, 3.6, 2.0, fill, line=None if fill == NAVY else LINE)
    text(s, t, x + 0.2, 2.62, 3.2, 0.8, size=17, bold=True, color=col, align=PP_ALIGN.CENTER, spacing=22)
    text(s, d, x + 0.2, 3.48, 3.2, 0.7, size=12, color=sub, align=PP_ALIGN.CENTER, spacing=15)
text(s, "→", 4.25, 3.05, 0.55, 0.5, size=24, color=ORANGE, align=PP_ALIGN.CENTER)
text(s, "→", 8.5, 3.05, 0.55, 0.5, size=24, color=ORANGE, align=PP_ALIGN.CENTER)
text(s, "забор пачками,\nкаждые 5 минут", 3.35, 4.45, 2.4, 0.6, size=10.5, color=SLATE,
     align=PP_ALIGN.CENTER, spacing=13)
text(s, "выгрузка результатов\nс повторами при сбое связи", 7.6, 4.45, 2.4, 0.6, size=10.5,
     color=SLATE, align=PP_ALIGN.CENTER, spacing=13)
notes3 = [
    ("Разворачивание", "один сервер у вас: локально или в частном облаке нужной юрисдикции"),
    ("Никаких внешних сервисов", "ни один документ не покидает ваш контур"),
    ("Совместимость", "работает поверх существующего хранилища и старой учётной системы"),
]
for i, (t, d) in enumerate(notes3):
    x = 0.6 + i * 4.25
    card(s, x, 5.4, 3.85, 1.5, BG_SOFT)
    text(s, t, x + 0.25, 5.6, 3.35, 0.3, size=13.5, bold=True)
    text(s, d, x + 0.25, 5.94, 3.35, 0.85, size=11.5, color=SLATE, spacing=15)
notes(s, "Минимальный след: не заменяем инфраструктуру, а подключаемся к ней.")

# ─────────────────────────── 15. Качество распознавания ───────────────────────────
s = head("Мы не обещаем точность — мы её измеряем", "Качество распознавания")
cd = CategoryChartData()
cd.categories = ["Базовая\nнастройка", "Режим\nстраницы", "Без жёсткой\nбинаризации",
                 "Оптимальная\nсвязка", "+ фильтр\nязыка"]
cd.add_series("Доля ошибочных символов", (0.309, 0.296, 0.219, 0.202, 0.198))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, I(0.6), I(1.95), I(7.5), I(4.1), cd)
ch = gf.chart
ch.has_title = False
ch.has_legend = False
plot = ch.plots[0]
plot.gap_width = 60
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = "0.000"
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(11)
dl.font.bold = True
dl.font.color.rgb = INK
dl.font.name = FONT
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = ORANGE
ca = ch.category_axis
ca.has_major_gridlines = False
ca.tick_labels.font.size = Pt(10)
ca.tick_labels.font.color.rgb = SLATE
ca.tick_labels.font.name = FONT
va = ch.value_axis
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = LINE
va.minimum_scale = 0.0
va.maximum_scale = 0.36
va.major_unit = 0.1
va.tick_labels.number_format = "0.00"
va.tick_labels.number_format_is_linked = False
va.tick_labels.font.size = Pt(9)
va.tick_labels.font.color.rgb = SLATE
va.tick_labels.font.name = FONT
text(s, "Доля ошибочных символов на реальном листе фонда (стенограмма 1931 г.); меньше — лучше",
     0.6, 6.15, 7.5, 0.5, size=11, color=SLATE, italic=True, spacing=14)
bullets(s, [
    "Эталонный набор: страницы вашего фонда с выверенной вручную расшифровкой",
    "Настройка принимается, только если ошибка уменьшилась — не «на глаз»",
    "Сделано: подавление цветных правок и печатей, выравнивание перекоса, "
    "словарь эпохи, отсев ложных букв",
    "−36% ошибок на этом листе за один вечер работы",
], 8.4, 2.05, 4.3, 3.3, size=12.5, spacing=17)
card(s, 8.4, 5.35, 4.3, 1.55, AMBER_BG)
text(s, "Что даст дальнейший рост", 8.65, 5.53, 3.8, 0.3, size=13, bold=True, color=AMBER_TX)
text(s, "Качественный скан 300 dpi вместо фотографии и видеокарта на сервере: "
        "под них у системы уже есть готовый режим работы.",
     8.65, 5.88, 3.8, 0.9, size=11.5, color=AMBER_TX, spacing=15)
notes(s, "Честный слайд: метрика и путь роста. Архивисты ценят измеримость.")

# ─────────────────────────── 16. Сравнение ───────────────────────────
s = head("Что вы получаете сверх показанного ранее", "Сравнение")
rows = [
    ("Ассистент с ответами по фонду и ссылками", "нет", "работает"),
    ("Настройка обычными словами", "обещано", "работает"),
    ("Своды персоналий без дублей", "дубли остаются", "сводятся автоматически"),
    ("Цифры на дашборде", "расходятся между экранами", "один источник"),
    ("Происхождение и уверенность данных", "нет", "по каждому документу"),
    ("Измеримое качество распознавания", "не показано", "эталон и метрика"),
    ("Модели и код", "закрытые", "открытые, заменяемые"),
]
y0, rh = 1.95, 0.63
text(s, "Возможность", 0.75, y0, 6.0, 0.32, size=12, bold=True, color=SLATE)
text(s, "Показанная ранее система", 6.9, y0, 3.0, 0.32, size=12, bold=True, color=SLATE,
     align=PP_ALIGN.CENTER)
text(s, "ArchivAI", 10.1, y0, 2.6, 0.32, size=12, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
for i, (a, b, c) in enumerate(rows):
    y = y0 + 0.5 + i * rh
    if i % 2 == 0:
        card(s, 0.6, y - 0.06, 12.1, rh - 0.04, BG_SOFT, radius=0.04)
    text(s, a, 0.75, y + 0.05, 6.0, 0.45, size=13)
    text(s, b, 6.9, y + 0.05, 3.0, 0.45, size=12, color=SLATE, align=PP_ALIGN.CENTER)
    text(s, c, 10.1, y + 0.05, 2.6, 0.45, size=12, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
notes(s, "Конкурента по имени в раздатке не называем — «показанная ранее система».")

# ─────────────────────────── 17. Технологии ───────────────────────────
s = head("Открытая, проверяемая, ваша", "Технологическая основа")
tech = [
    ("Открытые модели", "Распознавание и извлечение — на свободных моделях. Никаких подписок и лимитов, "
                        "никакой передачи документов наружу."),
    ("Прозрачный конвейер", "Каждый этап — видимый процесс, который можно посмотреть, повторить "
                            "и изменить под ваш регламент."),
    ("Надёжное хранение", "Документы, тексты, сущности и поисковый индекс — в одной базе. "
                          "Резервные копии — вашими стандартными средствами."),
    ("Рост под нагрузку", "На сервере с видеокартой точность и скорость растут без изменения системы: "
                          "меняется только модель."),
]
for i, (t, d) in enumerate(tech):
    x = 0.6 + (i % 2) * 6.25
    y = 1.95 + (i // 2) * 2.3
    card(s, x, y, 5.85, 2.0, BG_SOFT)
    circle(s, x + 0.3, y + 0.32, 0.55, NAVY, str(i + 1), size=15)
    text(s, t, x + 1.05, y + 0.35, 4.5, 0.35, size=16, bold=True)
    text(s, d, x + 1.05, y + 0.78, 4.5, 1.1, size=12, color=SLATE, spacing=16)
text(s, "Разворачивание: на вашем сервере, в вашей сети, с вашими правилами доступа.",
     0.6, 6.7, 12.1, 0.4, size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
notes(s, "Без жаргона: открыто, воспроизводимо и остаётся у клиента.")

# ─────────────────────────── 18. Дальше ───────────────────────────
s = head("Как начинаем", "Следующие шаги", dark=True)
steps = [
    ("1–2 недели", "Пилот на вашем фонде",
     "Берём 200–500 сканов, настраиваем справочник типов под ваши регламенты, "
     "показываем результат на ваших документах."),
    ("3–4 недели", "Внедрение в контуре",
     "Установка на ваш сервер, подключение хранилища и учётной системы, "
     "единая авторизация, обучение сотрудников."),
    ("дальше", "Развитие",
     "Дообучение на ваших проверенных документах, вырезка подписей и печатей, "
     "датчики хранилища, отчётность под ваши формы."),
]
for i, (badge, title, desc) in enumerate(steps):
    y = 2.0 + i * 1.55
    card(s, 0.6, y, 2.2, 0.62, ORANGE)
    text(s, badge, 0.6, y + 0.15, 2.2, 0.35, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, title, 3.1, y + 0.02, 9.5, 0.4, size=17, bold=True, color=WHITE)
    text(s, desc, 3.1, y + 0.47, 9.5, 0.8, size=12.5, color=MUTED, spacing=17)
text(s, "Система уже работает: сегодня мы показываем не прототип, а готовый контур.",
     0.6, 6.7, 12.1, 0.4, size=14, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
notes(s, "Закрытие: предложить пилот на их собственных документах.")

prs.save(OUT)
print("saved:", OUT, "| slides:", len(prs.slides.__iter__.__self__._sldIdLst))
