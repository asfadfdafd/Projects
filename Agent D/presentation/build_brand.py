#!/usr/bin/env python3
"""ArchivAI — презентация для клиента-архива в фирменном стиле JI Studio.

Бренд (brandBook.pdf):
  Deep Navy  000F85 — ядро, полноформатные плашки, первичные действия (не для текста)
  Steel Blue 5980A6 — акцент «рой», интерфейсные акценты
  Ink        101F20 — текст и линейки
  Ground     F2F2F3 — фон по умолчанию
  Paper      FFFFFF — выворотка, текст на navy
  Пропорция: Ground 50 · Paper 20 · Steel 15 · Navy 10 · Ink 5
  Заголовки — Barlow Condensed 600 (uppercase), подзаголовки — Barlow Condensed 400,
  текст — Barlow 400/500/700. Марка — только из поставленных локапов, без поворота.

Сборка:  python3 build_brand.py [out.pptx]
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches as I, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "img")
OUT = sys.argv[1] if len(sys.argv) > 1 else "ArchivAI_JIStudio.pptx"

# ── Бренд ────────────────────────────────────────────────────────────────────
NAVY = C(0x00, 0x0F, 0x85)      # Deep Navy — ядро
STEEL = C(0x59, 0x80, 0xA6)     # Steel Blue — акцент
STEEL_D = C(0x3C, 0x5C, 0x7A)   # глубокий шаг стали — для абзацного акцентного текста
INK = C(0x10, 0x1F, 0x20)       # текст
GROUND = C(0xF2, 0xF2, 0xF3)    # фон
PAPER = C(0xFF, 0xFF, 0xFF)
RULE = C(0xD8, 0xD8, 0xDC)      # хайрлайн
MUTE = C(0x6B, 0x70, 0x72)      # вторичный текст (Ink 60%)
MUTE_ON_NAVY = C(0xB9, 0xC4, 0xD6)
H = "Barlow Condensed"          # заголовки
B = "Barlow"                    # текст

prs = Presentation()
prs.slide_width, prs.slide_height = I(13.333), I(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 18
_page = {"n": 0}


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, s, x, y, w, h, size=14, color=INK, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, font=B,
         caps_track=False):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = Pt(spacing)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = font
        if caps_track:                       # трекинг для набора капителью
            r.font._rPr.set("spc", "120")
    return tb


def bullets(slide, items, x, y, w, h, size=14, color=INK, spacing=20, gap=9,
            marker_color=STEEL):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(spacing)
        p.space_after = Pt(gap)
        r0 = p.add_run()
        r0.text = "— "
        r0.font.size = Pt(size)
        r0.font.color.rgb = marker_color
        r0.font.name = B
        r = p.add_run()
        r.text = it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = B
    return tb


def card(slide, x, y, w, h, fill=PAPER, line=RULE):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def outline(slide, x, y, w, h, line=STEEL):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    sh.fill.background()
    sh.line.color.rgb = line
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def hairline(slide, x, y, w, color=RULE, weight=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), Emu(9525))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def num_box(slide, x, y, d, label, fill=NAVY, color=PAPER, size=13):
    """Квадратный номерной знак — в духе нумерации разделов брендбука."""
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(d), I(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = H
    return sh


def crossmark(slide, x, y, dark=False, size=0.11):
    """Крестик приводки — мотив брендбука."""
    col = C(0x4A, 0x55, 0x7A) if dark else RULE
    hairline(slide, x - size, y, size * 2, color=col, weight=0.5)
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(I(x).emu)), I(y - size),
                                Emu(9525), I(size * 2))
    sh.fill.solid()
    sh.fill.fore_color.rgb = col
    sh.line.fill.background()
    sh.shadow.inherit = False


def pic(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(os.path.join(IMG, name), I(x), I(y),
                                    I(w) if w else None, I(h) if h else None)


def header(slide, dark=False, fig=None):
    """Верхний колонтитул: марка · «JI STUDIO — ARCHIVAI» · номер страницы."""
    _page["n"] += 1
    mark = "mark_light.png" if dark else "mark_dark.png"
    slide.shapes.add_picture(os.path.join(IMG, mark), I(0.6), I(0.31), height=I(0.17))
    text(slide, "JI STUDIO  —  ARCHIVAI", 0.95, 0.3, 6, 0.24, size=9,
         color=MUTE_ON_NAVY if dark else MUTE, font=H, caps_track=True)
    right = f"{_page['n']:02d} / {TOTAL}"
    if fig:
        text(slide, fig, 8.4, 0.3, 3.0, 0.24, size=9, color=MUTE_ON_NAVY if dark else MUTE,
             font=H, align=PP_ALIGN.RIGHT, caps_track=True)
    text(slide, right, 11.6, 0.3, 1.13, 0.24, size=9, color=MUTE_ON_NAVY if dark else MUTE,
         font=H, align=PP_ALIGN.RIGHT, caps_track=True)
    hairline(slide, 0.6, 0.62, 12.13, color=C(0x2A, 0x36, 0x7A) if dark else RULE)


def head(title, kicker=None, dark=False, fig=None):
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY if dark else GROUND)
    header(s, dark=dark, fig=fig)
    if kicker:
        num_box(s, 0.6, 0.95, 0.3, f"{_page['n'] - 1:02d}",
                fill=STEEL if dark else NAVY, size=11)
        text(s, kicker.upper(), 1.05, 0.98, 8, 0.28, size=13, bold=True,
             color=PAPER if dark else INK, font=H, caps_track=True)
    text(s, title.upper(), 0.6, 1.38, 12.13, 0.72, size=30, bold=True,
         color=PAPER if dark else INK, font=H)
    return s


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


# ═══════════════════ 1. Титул ═══════════════════
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
_page["n"] += 1
pic(s, "logo_light.png", 0.6, 0.55, w=1.9)
text(s, "APPLIED AI  ·  AUTONOMOUS SYSTEMS", 0.62, 1.05, 6, 0.24, size=9,
     color=MUTE_ON_NAVY, font=H, caps_track=True)
text(s, "ARCHIVAI", 0.6, 2.35, 9, 1.25, size=76, bold=True, color=PAPER, font=H)
text(s, "Искусственный интеллект\nдля исторического архива", 0.6, 3.7, 8.4, 1.1,
     size=23, color=C(0xD5, 0xDC, 0xE8), font=H, spacing=30)
hairline(s, 0.6, 5.0, 5.2, color=STEEL)
text(s, "Сканы → распознанный текст, структурированные поля, связи между людьми и организациями,\n"
        "семантический поиск и проверка человеком. На открытых моделях, на вашем оборудовании.",
     0.6, 5.25, 8.0, 0.9, size=13, color=MUTE_ON_NAVY, spacing=19)
meta = [("КЛИЕНТ", "Государственный архив"), ("ПРОДУКТ", "ArchivAI · v1"),
        ("РАЗВЁРТЫВАНИЕ", "On-premise"), ("ДАТА", "07 · 2026")]
for i, (k, v) in enumerate(meta):
    x = 9.35
    y = 2.5 + i * 0.92
    text(s, k, x, y, 3.4, 0.22, size=9, color=STEEL, font=H, caps_track=True)
    text(s, v, x, y + 0.26, 3.4, 0.35, size=15, color=PAPER, font=H)
    if i < 3:
        hairline(s, x, y + 0.72, 3.38, color=C(0x2A, 0x36, 0x7A))
notes(s, "Титул. Тезис: это работающая система, а не концепт — дальше живые экраны.")

# ═══════════════════ 2. Задача ═══════════════════
s = head("Что мешает работать с фондом сегодня", "Задача", fig="FIG. 01 — КОНТЕКСТ")
items = [
    ("Документы не ищутся", "Скан — это картинка. Найти дело можно только по описи и памяти сотрудника."),
    ("Связи не видны", "Кто с кем переписывался, где упоминается человек — собирается вручную, месяцами."),
    ("Ручная опись", "Тип, дата, отправитель, получатель, персоналии — всё вносится руками."),
    ("Данные должны остаться у вас", "Облачные сервисы и передача фонда наружу — неприемлемы."),
]
for i, (t, d) in enumerate(items):
    y = 2.35 + i * 1.18
    card(s, 0.6, y, 7.3, 1.0)
    num_box(s, 0.6, y, 0.42, f"{i + 1:02d}", size=12)
    text(s, t, 1.25, y + 0.15, 6.4, 0.3, size=15, bold=True, font=H)
    text(s, d, 1.25, y + 0.5, 6.5, 0.45, size=11.5, color=MUTE, spacing=14)
pic(s, "scan_source.png", 8.4, 2.35, w=2.95)
text(s, "Типовой лист фонда: машинопись, правки от руки,\nпечать, пожелтевшая бумага",
     8.4, 6.85, 4.0, 0.5, size=9.5, color=MUTE, spacing=12)
crossmark(s, 12.6, 2.35)
notes(s, "Боли клиента. Последний пункт подводит к on-prem и открытым моделям.")

# ═══════════════════ 3. Как это работает ═══════════════════
s = head("От скана до знания — за один проход", "Как это работает", fig="FIG. 02 — КОНВЕЙЕР")
steps = [
    ("Приём", "загрузка или забор\nиз вашего хранилища"),
    ("Распознавание", "печатный и рукописный\nтекст, RU + KZ"),
    ("Извлечение", "тип, дата, отправитель,\nполучатель, резюме"),
    ("Сущности", "люди, организации,\nместа, даты, количества"),
    ("Индексация", "поиск по смыслу,\nа не по буквам"),
    ("Проверка", "правка человеком\nи статус «проверено»"),
]
for i, (t, d) in enumerate(steps):
    x = 0.6 + i * 2.06
    card(s, x, 2.5, 1.86, 2.35)
    num_box(s, x, 2.5, 0.42, f"{i + 1:02d}", size=12)
    text(s, t, x + 0.12, 3.15, 1.7, 0.3, size=12.5, bold=True, font=H)
    text(s, d, x + 0.16, 3.5, 1.62, 1.0, size=10.5, color=MUTE, spacing=13.5)
    if i < 5:
        text(s, "→", x + 1.86, 3.45, 0.2, 0.3, size=13, color=STEEL, align=PP_ALIGN.CENTER)
card(s, 0.6, 5.3, 12.13, 1.45, fill=PAPER)
text(s, "Каждый этап — отдельный видимый процесс", 0.95, 5.52, 11.4, 0.32, size=16, bold=True, font=H)
text(s, "Модель распознавания, правила извлечения, набор типов документов можно заменить или дополнить "
        "без переписывания системы. Ничего не «зашито» в чёрный ящик: любой шаг открывается, "
        "повторяется и настраивается под ваш фонд.",
     0.95, 5.9, 11.4, 0.75, size=12, color=MUTE, spacing=16)
notes(s, "Ключевая мысль: конвейер прозрачен и расширяем, модели меняются как детали.")

# ═══════════════════ 4. Приём документов ═══════════════════
s = head("Загрузка и забор из хранилища", "Приём документов", fig="FIG. 03 — ЭКРАН")
card(s, 0.55, 2.3, 7.3, 4.25)
pic(s, "upload.png", 0.7, 2.45, w=7.0)
bullets(s, [
    "Перетащите сканы в окно — или соберите многостраничное дело из нескольких файлов",
    "«Забрать из хранилища»: система сама берёт новые файлы — 50–100+ документов за раз",
    "Автоматический забор каждые 5 минут: положили сканы в папку — они появились в архиве обработанными",
    "Повторная заливка не создаёт дублей: уже обработанные файлы пропускаются",
], 8.35, 2.45, 4.4, 4.0, size=13.5)
notes(s, "Связка с их файловым хранилищем: системе нужны только хранилище и база.")

# ═══════════════════ 5. Архив ═══════════════════
s = head("Весь фонд — в одной таблице", "Архив документов", fig="FIG. 04 — ЭКРАН")
card(s, 0.55, 2.3, 12.23, 3.75)
pic(s, "documents.png", 0.7, 2.45, w=11.93)
cols = [
    ("Статусы обработки", "обрабатывается · на проверке · проверено · ошибка"),
    ("Фильтры", "по типу, языку, статусу, дате документа"),
    ("Миниатюры", "виден сам скан, а не только имя файла"),
]
for i, (t, d) in enumerate(cols):
    x = 0.6 + i * 4.15
    text(s, t, x, 6.3, 3.9, 0.3, size=14, bold=True, font=H)
    text(s, d, x, 6.62, 3.9, 0.5, size=11.5, color=MUTE, spacing=14)
notes(s, "Статус обработки виден сразу, ошибки не теряются.")

# ═══════════════════ 6. Карточка документа ═══════════════════
s = head("Скан слева — извлечённые данные справа", "Карточка документа", fig="FIG. 05 — ЭКРАН")
card(s, 0.55, 2.3, 5.75, 4.75)
pic(s, "doc_detail.png", 0.7, 2.42, h=4.5)
rows = [
    ("Тип, дата, язык", "классификация по вашему справочнику типов"),
    ("Отправитель и получатель", "с должностью и ролью, если они есть в тексте"),
    ("Резюме документа", "2–4 предложения на русском: суть дела без чтения листа"),
    ("Сущности по 5 категориям", "люди · организации · места · даты · количества"),
    ("Полный текст", "постранично и целиком, с источником и уверенностью"),
]
for i, (t, d) in enumerate(rows):
    y = 2.42 + i * 0.93
    hairline(s, 6.75, y - 0.08, 6.0)
    text(s, t, 6.75, y + 0.05, 5.9, 0.3, size=14.5, bold=True, font=H)
    text(s, d, 6.75, y + 0.38, 5.9, 0.45, size=11.5, color=MUTE, spacing=14)
notes(s, "Резюме и поля извлекаются автоматически — экономия часов описания.")

# ═══════════════════ 7. Человек в контуре ═══════════════════
s = head("Последнее слово — за архивистом", "Проверка и доверие", fig="FIG. 06 — ЭКРАН")
card(s, 6.75, 2.3, 6.03, 4.75)
pic(s, "doc_verify.png", 6.88, 2.42, h=4.5)
bullets(s, [
    "Любую сущность можно исправить прямо в карточке: опечатка в фамилии, пропущенная организация",
    "Кнопка «Проверено» переводит документ в доверенный статус",
    "Каждая правка записывается в журнал: кто, что и когда изменил",
    "Проверенные документы защищены: массовая переобработка их не трогает",
    "Исправления накапливаются как эталон — на них система дообучается под ваш фонд",
], 0.6, 2.42, 5.75, 3.5, size=13.5)
card(s, 0.6, 5.85, 5.75, 1.2, fill=PAPER, line=STEEL)
text(s, "ИИ ПРЕДЛАГАЕТ — ЧЕЛОВЕК УТВЕРЖДАЕТ", 0.85, 6.05, 5.3, 0.3, size=14, bold=True,
     color=NAVY, font=H)
text(s, "Ни одна запись не попадает в фонд как «истина» без подтверждения сотрудника.",
     0.85, 6.4, 5.3, 0.5, size=11.5, color=STEEL_D, spacing=14)
notes(s, "Ответ на главный страх архива: ИИ ничего не портит и не решает за людей.")

# ═══════════════════ 8. Поиск ═══════════════════
s = head("Поиск по смыслу, а не по буквам", "Семантический поиск", fig="FIG. 07 — ЭКРАН")
card(s, 0.55, 2.3, 7.35, 2.85)
pic(s, "search_results.png", 0.7, 2.42, w=7.05)
card(s, 0.55, 5.35, 7.35, 1.7)
pic(s, "search_recent.png", 0.7, 5.5, w=7.05)
bullets(s, [
    "Запрос обычными словами: «что-нибудь о кулачестве» находит приказ о ликвидации кулачества",
    "Работает на русском и казахском, в том числе на старых написаниях",
    "У каждого результата — оценка релевантности; слабые совпадения отсекаются",
    "История запросов у каждого сотрудника своя: можно вернуться к прошлому поиску",
], 8.35, 2.45, 4.4, 4.3, size=13.5)
notes(s, "Разница с поиском по подстроке — это клиент просил явно.")

# ═══════════════════ 9. Ask ArchivAI ═══════════════════
s = head("Вопрос архиву — ответ со ссылками", "Ассистент-исследователь",
         dark=True, fig="FIG. 08 — ОТЛИЧИЕ")
pic(s, "chat.png", 0.6, 2.45, w=6.85)
bullets(s, [
    "Обычный вопрос вместо запроса: «Кто и кому писал о юбилейных мероприятиях?»",
    "Ответ составляется только по вашим документам — каждая мысль подкреплена ссылкой",
    "Клик по ссылке открывает документ-первоисточник",
    "Если ответа в фонде нет, ассистент прямо об этом говорит и не выдумывает",
    "Диалог помнит контекст: вопрос можно уточнять дальше",
], 7.9, 2.5, 4.85, 3.6, size=13.5, color=C(0xE8, 0xEC, 0xF4), marker_color=STEEL)
outline(s, 7.9, 6.15, 4.85, 0.85, line=STEEL)
text(s, "У показанной ранее системы этого нет:\nеё поиск возвращает только список документов",
     8.15, 6.3, 4.4, 0.6, size=11.5, bold=True, color=PAPER, spacing=15)
notes(s, "Дифференциатор №1. У конкурента это только пожелание клиента, у нас работает.")

# ═══════════════════ 10. Агент ═══════════════════
s = head("Настройка системы обычными словами", "Агент · флагман", dark=True,
         fig="FIG. 09 — ОТЛИЧИЕ")
pic(s, "agent.png", 0.6, 2.45, w=6.95)
text(s, "СКАЖИТЕ — И СИСТЕМА ПЕРЕСТРОИТСЯ:", 7.95, 2.45, 4.8, 0.3, size=13, bold=True,
     color=STEEL, font=H, caps_track=True)
phrases = [
    "«Добавь тип документа „протокол“\nи переклассифицируй архив»",
    "«Добавь виджет: динамика приказов по годам»",
    "«Сделай поиск более смысловым»",
    "«Слей дубликаты персоналий»",
]
for i, p_ in enumerate(phrases):
    y = 2.9 + i * 0.78
    outline(s, 7.95, y, 4.8, 0.66, line=C(0x2A, 0x36, 0x7A))
    text(s, p_, 8.15, y + 0.1, 4.5, 0.5, size=11, color=C(0xE8, 0xEC, 0xF4), spacing=13)
outline(s, 7.95, 6.1, 4.8, 0.95, line=STEEL)
text(s, "Безопасно по построению: модель не имеет доступа к базе — она лишь предлагает изменение. "
        "База проверяет его по правилам, применяет и записывает в журнал.",
     8.15, 6.22, 4.45, 0.75, size=10.5, color=C(0xD5, 0xDC, 0xE8), spacing=13)
notes(s, "Дифференциатор №2. Конкурент это только обещал — «поэтому продукт и назван Agent D».")

# ═══════════════════ 11. Сущности и граф ═══════════════════
s = head("Связи людей, организаций и мест", "Сущности и связи",
         fig="FIG. 10 — ЭКРАН")
card(s, 0.55, 2.3, 7.7, 4.6)
pic(s, "entities.png", 0.68, 2.42, w=7.44)
bullets(s, [
    "Одна персона — одна карточка: разные написания фамилии сводятся автоматически, "
    "старые сохраняются как псевдонимы",
    "Граф показывает, кто встречается в одних и тех же делах — связи, которые вручную не увидеть",
    "От сущности — переход ко всем её документам",
    "Выгрузка списка сущностей в CSV для отчётов и описей",
], 8.6, 2.45, 4.15, 4.3, size=13.5)
notes(s, "Наше преимущество: у конкурента дубликаты «Абай» / «Абай Кунанбаев» живут отдельно.")

# ═══════════════════ 12. Дашборд + IoT ═══════════════════
s = head("Состояние фонда и хранилища", "Аналитика", fig="FIG. 11 — ЭКРАН")
card(s, 0.55, 2.3, 8.4, 4.6)
pic(s, "dashboard.png", 0.68, 2.42, w=8.14)
bullets(s, [
    "Сколько документов обработано, сколько персоналий и организаций найдено",
    "Распределение по типам, языкам, хронологии фонда",
    "Все цифры считаются из одного источника — они не расходятся между экранами",
    "Датчики хранилища: температура и влажность в реальном времени, с нормами и историей за сутки",
    "Любой ваш датчик шлёт данные на один адрес — виджет появляется сам",
], 9.25, 2.45, 3.5, 4.4, size=12, spacing=17)
notes(s, "IoT-виджет — прямой запрос клиента на демо конкурента.")

# ═══════════════════ 13. Комплаенс ═══════════════════
s = head("Кто, что и когда смотрел", "Учёт и контроль", fig="FIG. 12 — ЭКРАН")
card(s, 4.6, 2.3, 8.18, 4.75)
pic(s, "usage.png", 4.72, 2.42, h=4.5)
bullets(s, [
    "Каждый вход — с адресом и временем, включая неудачные попытки",
    "Каждый поиск и каждое открытие документа фиксируются",
    "Видно, какие дела востребованы и кто из сотрудников сейчас работает",
    "Права: администратор ведёт справочники, исследователь работает с фондом",
    "Готово к подключению вашей единой авторизации (SSO)",
], 0.6, 2.45, 3.75, 4.6, size=12, spacing=17)
notes(s, "Ответ на «compliance»: журнал, адреса, востребованность документов.")

# ═══════════════════ 14. Интеграция ═══════════════════
s = head("Нужны только ваше хранилище и база", "Интеграция",
         fig="FIG. 13 — СХЕМА")
boxes = [
    (0.6, "Ваше файловое\nхранилище", "сканы как есть", PAPER, INK, MUTE, RULE),
    (4.87, "ArchivAI", "распознавание, извлечение,\nпоиск, проверка", NAVY, PAPER, MUTE_ON_NAVY, None),
    (9.14, "Ваша база\nи учётная система", "структурированные данные", PAPER, INK, MUTE, RULE),
]
for x, t, d, fill, col, sub, ln in boxes:
    card(s, x, 2.75, 3.6, 1.95, fill=fill, line=ln)
    text(s, t, x + 0.25, 3.05, 3.1, 0.75, size=17, bold=True, color=col,
         align=PP_ALIGN.CENTER, font=H, spacing=21)
    text(s, d, x + 0.25, 3.9, 3.1, 0.6, size=11, color=sub, align=PP_ALIGN.CENTER, spacing=14)
text(s, "→", 4.27, 3.5, 0.5, 0.4, size=20, color=STEEL, align=PP_ALIGN.CENTER)
text(s, "→", 8.54, 3.5, 0.5, 0.4, size=20, color=STEEL, align=PP_ALIGN.CENTER)
text(s, "забор пачками,\nкаждые 5 минут", 3.3, 4.85, 2.4, 0.55, size=10, color=MUTE,
     align=PP_ALIGN.CENTER, spacing=12)
text(s, "выгрузка результатов\nс повторами при сбое связи", 7.6, 4.85, 2.4, 0.55, size=10,
     color=MUTE, align=PP_ALIGN.CENTER, spacing=12)
notes3 = [
    ("Разворачивание", "один сервер у вас: локально или в частном облаке нужной юрисдикции"),
    ("Никаких внешних сервисов", "ни один документ не покидает ваш контур"),
    ("Совместимость", "работает поверх существующего хранилища и старой учётной системы"),
]
for i, (t, d) in enumerate(notes3):
    x = 0.6 + i * 4.27
    hairline(s, x, 5.65, 3.85, color=NAVY, weight=1)
    text(s, t, x, 5.8, 3.85, 0.3, size=13.5, bold=True, font=H)
    text(s, d, x, 6.13, 3.85, 0.8, size=11.5, color=MUTE, spacing=14)
notes(s, "Минимальный след: не заменяем инфраструктуру, а подключаемся к ней.")

# ═══════════════════ 15. Качество распознавания ═══════════════════
s = head("Точность мы измеряем, а не обещаем", "Качество распознавания",
         fig="FIG. 14 — ИЗМЕРЕНИЕ")
cd = CategoryChartData()
cd.categories = ["Базовая\nнастройка", "Режим\nстраницы", "Без жёсткой\nбинаризации",
                 "Оптимальная\nсвязка", "+ фильтр\nязыка"]
cd.add_series("Доля ошибочных символов", (0.309, 0.296, 0.219, 0.202, 0.198))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, I(0.6), I(2.35), I(7.5), I(4.0), cd)
ch = gf.chart
ch.has_title = False
ch.has_legend = False
plot = ch.plots[0]
plot.gap_width = 70
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = "0.000"
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(11)
dl.font.bold = True
dl.font.color.rgb = INK
dl.font.name = H
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = NAVY
ca = ch.category_axis
ca.has_major_gridlines = False
ca.format.line.color.rgb = RULE
ca.tick_labels.font.size = Pt(10)
ca.tick_labels.font.color.rgb = MUTE
ca.tick_labels.font.name = B
va = ch.value_axis
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = RULE
va.minimum_scale = 0.0
va.maximum_scale = 0.36
va.major_unit = 0.1
va.tick_labels.number_format = "0.00"
va.tick_labels.number_format_is_linked = False
va.tick_labels.font.size = Pt(9)
va.tick_labels.font.color.rgb = MUTE
va.tick_labels.font.name = B
text(s, "Доля ошибочных символов на реальном листе фонда (стенограмма 1931 г.); меньше — лучше",
     0.6, 6.45, 7.5, 0.45, size=10.5, color=MUTE, italic=True, spacing=13)
bullets(s, [
    "Эталонный набор: страницы вашего фонда с выверенной вручную расшифровкой",
    "Настройка принимается, только если ошибка уменьшилась — не «на глаз»",
    "Сделано: подавление цветных правок и печатей, выравнивание перекоса, "
    "словарь эпохи, отсев ложных букв",
    "−36% ошибок на этом листе за один вечер работы",
], 8.4, 2.45, 4.35, 3.2, size=12, spacing=16)
card(s, 8.4, 5.6, 4.35, 1.45, fill=PAPER, line=STEEL)
text(s, "ЧТО ДАСТ ДАЛЬНЕЙШИЙ РОСТ", 8.65, 5.78, 3.9, 0.28, size=12.5, bold=True,
     color=NAVY, font=H)
text(s, "Качественный скан 300 dpi вместо фотографии и видеокарта на сервере: "
        "под них у системы уже есть готовый режим работы.",
     8.65, 6.12, 3.9, 0.85, size=11, color=STEEL_D, spacing=14)
notes(s, "Честный слайд: метрика и путь роста. Архивисты ценят измеримость.")

# ═══════════════════ 16. Сравнение ═══════════════════
s = head("Что вы получаете сверх показанного", "Сравнение", fig="FIG. 15 — СВОДКА")
rows = [
    ("Ассистент с ответами по фонду и ссылками", "нет", "работает"),
    ("Настройка обычными словами", "обещано", "работает"),
    ("Своды персоналий без дублей", "дубли остаются", "сводятся автоматически"),
    ("Цифры на дашборде", "расходятся между экранами", "один источник"),
    ("Происхождение и уверенность данных", "нет", "по каждому документу"),
    ("Измеримое качество распознавания", "не показано", "эталон и метрика"),
    ("Модели и код", "закрытые", "открытые, заменяемые"),
]
y0, rh = 2.45, 0.6
text(s, "ВОЗМОЖНОСТЬ", 0.62, y0, 6.0, 0.3, size=11, bold=True, color=MUTE, font=H, caps_track=True)
text(s, "ПОКАЗАННАЯ РАНЕЕ СИСТЕМА", 6.9, y0, 3.0, 0.3, size=11, bold=True, color=MUTE,
     font=H, align=PP_ALIGN.CENTER, caps_track=True)
text(s, "ARCHIVAI", 10.1, y0, 2.6, 0.3, size=11, bold=True, color=NAVY, font=H,
     align=PP_ALIGN.CENTER, caps_track=True)
hairline(s, 0.6, y0 + 0.34, 12.13, color=INK, weight=1)
for i, (a, b_, c_) in enumerate(rows):
    y = y0 + 0.5 + i * rh
    if i % 2 == 0:
        card(s, 0.6, y - 0.06, 12.13, rh - 0.04, fill=PAPER, line=None)
    text(s, a, 0.75, y + 0.04, 6.0, 0.42, size=12.5)
    text(s, b_, 6.9, y + 0.04, 3.0, 0.42, size=11.5, color=MUTE, align=PP_ALIGN.CENTER)
    text(s, c_, 10.1, y + 0.04, 2.6, 0.42, size=11.5, bold=True, color=NAVY,
         align=PP_ALIGN.CENTER, font=H)
notes(s, "Конкурента по имени в раздатке не называем — «показанная ранее система».")

# ═══════════════════ 17. Технологии ═══════════════════
s = head("Открытая, проверяемая, ваша", "Технологическая основа", fig="FIG. 16 — ОСНОВА")
tech = [
    ("Открытые модели", "Распознавание и извлечение — на свободных моделях. Никаких подписок "
                        "и лимитов, никакой передачи документов наружу."),
    ("Прозрачный конвейер", "Каждый этап — видимый процесс, который можно посмотреть, "
                            "повторить и изменить под ваш регламент."),
    ("Надёжное хранение", "Документы, тексты, сущности и поисковый индекс — в одной базе. "
                          "Резервные копии — вашими стандартными средствами."),
    ("Рост под нагрузку", "На сервере с видеокартой точность и скорость растут без изменения "
                          "системы: меняется только модель."),
]
for i, (t, d) in enumerate(tech):
    x = 0.6 + (i % 2) * 6.28
    y = 2.45 + (i // 2) * 2.2
    card(s, x, y, 5.85, 1.9)
    num_box(s, x, y, 0.42, f"{i + 1:02d}", size=12)
    text(s, t, x + 0.62, y + 0.12, 5.0, 0.32, size=16, bold=True, font=H)
    text(s, d, x + 0.62, y + 0.55, 5.0, 1.1, size=11.5, color=MUTE, spacing=15)
hairline(s, 0.6, 6.95, 12.13, color=NAVY, weight=1)
text(s, "Разворачивание: на вашем сервере, в вашей сети, с вашими правилами доступа.",
     0.6, 7.05, 12.13, 0.3, size=12.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER, font=H)
notes(s, "Без жаргона: открыто, воспроизводимо и остаётся у клиента.")

# ═══════════════════ 18. Дальше ═══════════════════
s = head("Как начинаем", "Следующие шаги", dark=True, fig="FIG. 17 — ПЛАН")
steps = [
    ("1–2 НЕДЕЛИ", "Пилот на вашем фонде",
     "Берём 200–500 сканов, настраиваем справочник типов под ваши регламенты, "
     "показываем результат на ваших документах."),
    ("3–4 НЕДЕЛИ", "Внедрение в контуре",
     "Установка на ваш сервер, подключение хранилища и учётной системы, "
     "единая авторизация, обучение сотрудников."),
    ("ДАЛЬШЕ", "Развитие",
     "Дообучение на ваших проверенных документах, вырезка подписей и печатей, "
     "датчики хранилища, отчётность под ваши формы."),
]
for i, (badge, title, desc) in enumerate(steps):
    y = 2.5 + i * 1.35
    hairline(s, 0.6, y - 0.12, 12.13, color=C(0x2A, 0x36, 0x7A))
    card(s, 0.6, y, 1.85, 0.42, fill=STEEL, line=None)
    text(s, badge, 0.6, y + 0.08, 1.85, 0.28, size=11, bold=True, color=PAPER,
         align=PP_ALIGN.CENTER, font=H, caps_track=True)
    text(s, title, 2.75, y - 0.02, 9.9, 0.35, size=18, bold=True, color=PAPER, font=H)
    text(s, desc, 2.75, y + 0.4, 9.9, 0.7, size=12, color=MUTE_ON_NAVY, spacing=16)
hairline(s, 0.6, 6.55, 12.13, color=STEEL)
text(s, "Система уже работает: сегодня мы показываем не прототип, а готовый контур.",
     0.6, 6.7, 8.0, 0.35, size=13.5, bold=True, color=PAPER, font=H)
pic(s, "logo_light.png", 11.05, 6.62, w=1.68)
notes(s, "Закрытие: предложить пилот на их собственных документах.")

prs.save(OUT)
print("saved:", OUT)
